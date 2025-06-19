from docx import Document

def extract_placeholders(path):
    doc = Document(path)
    placeholders = set()
    for para in doc.paragraphs:
        if "{{" in para.text:
            words = para.text.split()
            for word in words:
                if word.startswith("{{") and word.endswith("}}"):
                    placeholders.add(word.strip("{{}}"))
    return list(placeholders)
